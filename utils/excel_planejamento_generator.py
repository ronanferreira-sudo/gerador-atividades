import os
import re
import math
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime
from utils.texto_utils import limpar_numeracao_texto, limpar_numeracao_item


def calcular_altura_linha(texto_o_que, texto_estrategia="", texto_recursos="", texto_criterios="", texto_instrumentos=""):
    """
    Calcula a altura ideal da linha do Excel baseando-se no número de linhas reais
    e quebras automáticas de texto (wrap text) para garantir que nenhuma informação
    fique cortada ou oculta.
    """
    def estimar_linhas_bloco(texto, largura_chars):
        if not texto:
            return 0
        total = 0
        for linha in str(texto).split("\n"):
            linha_s = linha.strip()
            if not linha_s:
                total += 1
            else:
                total += max(1, math.ceil(len(linha_s) / max(largura_chars, 10)))
        return total

    # Larguras estimadas em caracteres úteis por coluna mesclada:
    # A:C (~50 caracteres)
    # D (~32 caracteres)
    # H:I (~16 caracteres)
    # M:O (~22 caracteres)
    # P (~18 caracteres)
    linhas_a = estimar_linhas_bloco(texto_o_que, 50)
    linhas_d = estimar_linhas_bloco(texto_estrategia, 32)
    linhas_h = estimar_linhas_bloco(texto_recursos, 16)
    linhas_m = estimar_linhas_bloco(texto_criterios, 22)
    linhas_p = estimar_linhas_bloco(texto_instrumentos, 18)

    max_linhas = max(linhas_a, linhas_d, linhas_h, linhas_m, linhas_p, 6)

    # Cada linha consome aproximadamente 16 pontos em Calibri 10, com margem interna
    altura_estimada = max_linhas * 16 + 25

    # Altura base mínima de 115 pt; aumenta dinamicamente conforme a quantidade de conteúdo
    return max(115, altura_estimada)


def gerar_excel_planejamento(
    dados_cabecalho,
    aulas,
    caminho_modelo="templates/modelo_planejamento_docente.xlsx",
    caminho_saida=None
):
    """
    Gera o arquivo Excel do Planejamento Docente a partir do modelo base,
    preenchendo cabeçalho e todas as aulas calculadas com a Taxonomia de Bloom,
    com informações limpas, organizadas e altura de linha adaptativa.
    """
    if not os.path.exists(caminho_modelo):
        raise FileNotFoundError(f"Arquivo modelo não encontrado: {caminho_modelo}")

    wb = openpyxl.load_workbook(caminho_modelo)
    ws = wb.active

    # ==========================================
    # 1. ATUALIZA O CABEÇALHO (Linhas 5 a 8)
    # ==========================================
    # C5: Unidade Escolar
    if dados_cabecalho.get("unidade_escolar"):
        ws["C5"] = str(dados_cabecalho["unidade_escolar"]).strip()

    # L5: Início e Fim
    data_ini = dados_cabecalho.get("data_inicio", "")
    data_fim = dados_cabecalho.get("data_fim", "")
    if isinstance(data_ini, (datetime,)):
        data_ini = data_ini.strftime("%d/%m/%Y")
    if isinstance(data_fim, (datetime,)):
        data_fim = data_fim.strftime("%d/%m/%Y")
    if data_ini and data_fim:
        ws["L5"] = f"{data_ini} a {data_fim}"
    elif data_ini:
        ws["L5"] = f"A partir de {data_ini}"

    # C6: Curso
    if dados_cabecalho.get("nome_curso"):
        ws["C6"] = str(dados_cabecalho["nome_curso"]).strip().upper()

    # L6: Modalidade
    modalidade = dados_cabecalho.get("modalidade", "HABILITAÇÃO TÉCNICA")
    ws["L6"] = f" {modalidade.strip().upper()} "

    # C7: Código da Turma
    turma = dados_cabecalho.get("codigo_turma", "TEC.2026.1")
    ws["C7"] = f" {turma.strip().upper()}"

    # L7: Carga Horária da U.C
    if dados_cabecalho.get("carga_horaria"):
        try:
            ws["L7"] = float(dados_cabecalho["carga_horaria"])
        except ValueError:
            ws["L7"] = str(dados_cabecalho["carga_horaria"])

    # C8: Unidade Curricular
    if dados_cabecalho.get("unidade_curricular"):
        ws["C8"] = str(dados_cabecalho["unidade_curricular"]).strip()

    # L8: Instrutor
    if dados_cabecalho.get("instrutor"):
        ws["L8"] = str(dados_cabecalho["instrutor"]).strip().upper()

    # ==========================================
    # 2. LIMPAR LINHAS ANTERIORES (Linha 13 em diante)
    # ==========================================
    # Remove mesclagens antigas das linhas de dados (>= 13)
    ranges_para_remover = [
        rng for rng in list(ws.merged_cells.ranges)
        if rng.min_row >= 13
    ]
    for rng in ranges_para_remover:
        ws.merged_cells.remove(rng)

    # Deleta as linhas anteriores a partir da linha 13
    if ws.max_row >= 13:
        ws.delete_rows(13, ws.max_row - 12)

    # Estilos padrão
    fonte_padrao = Font(name="Calibri", size=10, bold=False)
    fonte_link = Font(name="Calibri", size=9, bold=False, color="0000FF", underline="single")
    
    alin_esquerda_topo = Alignment(horizontal="left", vertical="top", wrap_text=True)
    alin_centro_topo = Alignment(horizontal="center", vertical="top", wrap_text=True)
    alin_centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alin_data = Alignment(horizontal="center", vertical="center", wrap_text=False)
    
    borda_fina = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000")
    )

    # ==========================================
    # 3. PREENCHER AS AULAS
    # ==========================================
    total_aulas = len(aulas)
    linha_inicio_dados = 13

    for idx, aula in enumerate(aulas):
        r = linha_inicio_dados + idx

        # Texto do bloco O que? (Capacidade / Conhecimento)
        # 100% sem números de tópicos e sem símbolos ou '?' antes das palavras
        cap_limpa = limpar_numeracao_texto(aula.get('capacidade', '')).strip()
        con_limpo = limpar_numeracao_texto(aula.get('conhecimento', ''), separador="\n").strip()
        
        # Títulos CAPACIDADE e CONHECIMENTOS centralizados visualmente,
        # enquanto o texto da capacidade e os conteúdos do plano continuam alinhados à esquerda
        titulo_cap = " " * 19 + "CAPACIDADE:"
        titulo_con = " " * 17 + "CONHECIMENTOS:"

        texto_o_que = (
            f"{titulo_cap}\n"
            f"{cap_limpa}\n\n"
            f"{titulo_con}\n"
            f"{con_limpo}"
        )
        ws.cell(row=r, column=1, value=texto_o_que)  # Col A

        # Col D: Como? Estratégia de Ensino
        estrategia_txt = str(aula.get("estrategia", "")).strip().upper()
        ws.cell(row=r, column=4, value=estrategia_txt)

        # Col E: Onde? Ambientes Pedagógicos
        ws.cell(row=r, column=5, value=str(aula.get("ambiente", "LABORATORIO T.I")).strip().upper())

        # Col H: Recursos Didáticos
        recursos_txt = str(aula.get("recursos", "NOTEBOOK\nTELA DIGITAL\nCOMPUTADORES")).strip().upper()
        ws.cell(row=r, column=8, value=recursos_txt)

        # Col J: Carga Horária
        ws.cell(row=r, column=10, value=str(aula.get("carga_horaria_str", "4 HS")).strip().upper())

        # Col M: Critérios de Avaliação
        criterios_txt = str(aula.get("criterios", "")).strip().upper()
        ws.cell(row=r, column=13, value=criterios_txt)

        # Col P: Instrumento de Avaliação
        instrumentos_txt = str(aula.get("instrumentos", "")).strip().upper()
        ws.cell(row=r, column=16, value=instrumentos_txt)

        # Col Q: Situação de Aprendizagem
        ws.cell(row=r, column=17, value=str(aula.get("situacao_aprendizagem", "https://recursosdidaticos.senai.br")).strip())

        # Datas (R e S)
        dt_ini = aula.get("data_inicio")
        dt_fim = aula.get("data_fim", dt_ini)

        cell_r = ws.cell(row=r, column=18)
        cell_s = ws.cell(row=r, column=19)

        if isinstance(dt_ini, datetime):
            cell_r.value = dt_ini
            cell_r.number_format = "DD/MM/YYYY"
        else:
            cell_r.value = str(dt_ini) if dt_ini else ""

        if isinstance(dt_fim, datetime):
            cell_s.value = dt_fim
            cell_s.number_format = "DD/MM/YYYY"
        else:
            cell_s.value = str(dt_fim) if dt_fim else ""

        # Calcula a altura necessária da linha dinamicamente para acomodar todo o conteúdo
        ws.row_dimensions[r].height = calcular_altura_linha(
            texto_o_que,
            estrategia_txt,
            recursos_txt,
            criterios_txt,
            instrumentos_txt
        )

        # Aplica estilos e bordas nas células de A a S da linha r
        for col_idx in range(1, 20):
            c_obj = ws.cell(row=r, column=col_idx)
            c_obj.font = fonte_padrao
            c_obj.border = borda_fina

            # Alinhamentos específicos para máxima organização visual
            if col_idx in (1, 2, 3):  # Colunas A:C (Capacidade / Conhecimento)
                c_obj.alignment = alin_esquerda_topo
            elif col_idx in (13, 14, 15):  # Colunas M:O (Critérios de Avaliação)
                c_obj.alignment = alin_esquerda_topo
            elif col_idx in (18, 19):  # Colunas R e S (Datas)
                c_obj.alignment = alin_data
            else:
                c_obj.alignment = alin_centro

        # Ajusta mesclagens da linha r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)   # A:C
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)   # E:G
        ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=9)   # H:I
        ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=12) # J:L
        ws.merge_cells(start_row=r, start_column=13, end_row=r, end_column=15) # M:O

    # Mescla a coluna Q (Situação de Aprendizagem) por toda a extensão se desejado
    if total_aulas > 1:
        primeira_linha = linha_inicio_dados
        ultima_linha = linha_inicio_dados + total_aulas - 1
        ws.merge_cells(start_row=primeira_linha, start_column=17, end_row=ultima_linha, end_column=17)
        ws.cell(row=primeira_linha, column=17).alignment = alin_centro

    # Define caminho de saída se não fornecido
    if not caminho_saida:
        os.makedirs("outputs", exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        nome_uc_safe = "".join(c for c in dados_cabecalho.get("unidade_curricular", "UC") if c.isalnum() or c in (" ", "_", "-")).strip()
        caminho_saida = f"outputs/Planejamento_Docente_{nome_uc_safe}_{timestamp}.xlsx"

    wb.save(caminho_saida)
    wb.close()
    return caminho_saida
